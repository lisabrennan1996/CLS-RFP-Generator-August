#!/usr/bin/env python3
"""Reconstruct the two Specimen-Management sections (Referral Lab Samples and
Storage Samples / LTS) directly from a previous, completed Central Lab RFP
.docx — which is built from the same template, so its data tables are real
Word tables (not OCR/positional reconstruction). Returns
{row_label: {column: value}} for each section, matching the shape
`populate_rfp.py::fill_spec()` already consumes.
"""
from __future__ import annotations

import docx


def _distinct(row):
    """De-duplicate cells that share the same underlying <w:tc> (merged cells)."""
    out, seen = [], set()
    for c in row.cells:
        if id(c._tc) not in seen:
            seen.add(id(c._tc))
            out.append(c)
    return out


def _find_data_table(doc, keyword: str, min_rows: int = 5):
    """Same header-keyword + row-count heuristic populate_rfp.py uses to
    locate template tables — the previous RFP shares the same template."""
    candidates = []
    for i, t in enumerate(doc.tables):
        try:
            header = t.rows[0].cells[0].text
        except (IndexError, AttributeError):
            continue
        if keyword.lower() in header.lower() and len(t.rows) >= min_rows:
            candidates.append((i, len(t.rows), len(t.columns)))
    if not candidates:
        return None
    return doc.tables[max(candidates, key=lambda x: (x[1], x[2]))[0]]


def _read_table(table, columns: list[str]) -> dict[str, dict[str, str]]:
    """Read a data table into {row_label: {column: value}}.

    `columns` is the list of target column names to extract (substring
    matched against the table's own header row, case-insensitive) — mirrors
    how `fill_spec()` locates columns when writing.
    """
    header_cells = _distinct(table.rows[0])
    col_idx = {}
    for col_name in columns:
        for i, c in enumerate(header_cells):
            if col_name.lower() in c.text.lower():
                col_idx[col_name] = i
                break

    out: dict[str, dict[str, str]] = {}
    for row in table.rows[1:]:
        cells = _distinct(row)
        if not cells:
            continue
        label = cells[0].text.strip()
        if not label:
            continue
        values = {}
        for col_name, idx in col_idx.items():
            if idx < len(cells):
                val = cells[idx].text.strip()
                if val:
                    values[col_name] = val
        if values:
            out[label] = values
    return out


def build(previous_rfp_path: str = '') -> tuple[dict, dict]:
    """Return (referral, storage) dicts, or ({}, {}) if no previous RFP is
    given or its tables can't be located."""
    if not previous_rfp_path:
        return {}, {}
    try:
        doc = docx.Document(previous_rfp_path)
    except Exception:
        return {}, {}

    referral_table = _find_data_table(doc, 'REFERRAL LAB')
    referral = _read_table(referral_table, ['LTS PK']) if referral_table is not None else {}

    storage_wide = _find_data_table(doc, 'STORAGE SAMPLES', min_rows=10)
    storage = _read_table(storage_wide, ['LTS DNA', 'LTS Serum', 'LTS Plasma']) if storage_wide is not None else {}

    # The RNA column lives in a second, narrower "STORAGE SAMPLES" table.
    storage_candidates = [
        t for t in doc.tables
        if t is not storage_wide
        and t.rows and 'storage samples' in t.rows[0].cells[0].text.lower()
        and len(t.rows) >= 10
    ]
    if storage_candidates:
        rna_table = storage_candidates[0]
        rna_data = _read_table(rna_table, ['LTS RNA'])
        for label, cols in rna_data.items():
            storage.setdefault(label, {}).update(cols)

    return referral, storage


if __name__ == '__main__':
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else ''
    ref, sto = build(path)
    print('=== REFERRAL (LTS PK) ===')
    for lab, c in ref.items():
        print(f'  {lab:42} -> {c.get("LTS PK", "")!r}')
    print('\n=== STORAGE (DNA | Serum | Plasma | RNA) ===')
    for lab, c in sto.items():
        print(f'  {lab:42} | {c.get("LTS DNA",""):20} | {c.get("LTS Serum",""):14} | '
              f'{c.get("LTS Plasma",""):12} | {c.get("LTS RNA","")}')
